"""Document parsing for PDF, DOCX and plain text.

Unlike a typical resume parser, this one keeps **span-level visual metadata**
(colour, font size, style flags) alongside the text. That is not decoration: it
is the evidence base for the invisible-text signal in
:mod:`talentlens.antigaming`. A parser that returns only a string makes
white-on-white keyword stuffing undetectable by construction, so the metadata is
captured here at the only point in the pipeline where it still exists.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from .schemas import ParsedDocument, TextSpan

try:  # PyMuPDF renamed its module; support both spellings.
    import pymupdf as _fitz
except ImportError:  # pragma: no cover - depends on installed version
    import fitz as _fitz  # type: ignore[no-redef]


class ParsingError(RuntimeError):
    """Raised when a document cannot be read at all."""


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}

_MULTI_NEWLINE = re.compile(r"\n{3,}")
_TRAILING_SPACE = re.compile(r"[ \t]+\n")

#: A full-page vector rectangle must cover at least this fraction of the page
#: before we treat its fill as the page background rather than as a decorative
#: box. Sidebars and header bands are common and must not be mistaken for it.
_BACKGROUND_AREA_RATIO = 0.9


def _tidy(text: str) -> str:
    text = _TRAILING_SPACE.sub("\n", text)
    return _MULTI_NEWLINE.sub("\n\n", text).strip()


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------


def _page_background(page) -> int:
    """Best-effort dominant background colour of a page, packed 0xRRGGBB.

    Defaults to white, which is correct for effectively every real resume. We
    only override it when a near-full-page filled rectangle is present, because
    that is the case where "text the same colour as the background" means
    something other than white.
    """
    try:
        page_area = abs(page.rect.get_area())
        if page_area <= 0:
            return 0xFFFFFF
        for drawing in page.get_drawings():
            fill = drawing.get("fill")
            if fill is None:
                continue
            rect = drawing.get("rect")
            if rect is None or abs(rect.get_area()) < page_area * _BACKGROUND_AREA_RATIO:
                continue
            r, g, b = (int(round(c * 255)) for c in fill[:3])
            return (r << 16) | (g << 8) | b
    except Exception:
        # Drawing extraction is best-effort; a malformed content stream must
        # never take down a screening run.
        return 0xFFFFFF
    return 0xFFFFFF


def _parse_pdf(data: bytes, path: str) -> ParsedDocument:
    try:
        doc = _fitz.open(stream=data, filetype="pdf")
    except Exception as exc:  # pragma: no cover - corrupt file path
        raise ParsingError(f"could not open PDF {path}: {exc}") from exc

    spans: list[TextSpan] = []
    backgrounds: list[int] = []
    page_texts: list[str] = []

    with doc:
        for page_number, page in enumerate(doc):
            backgrounds.append(_page_background(page))
            page_lines: list[str] = []
            layout = page.get_text("dict")
            for block in layout.get("blocks", ()):
                if block.get("type") != 0:  # 0 == text block; 1 == image
                    continue
                for line in block.get("lines", ()):
                    line_parts: list[str] = []
                    for span in line.get("spans", ()):
                        raw = span.get("text", "")
                        if not raw.strip():
                            continue
                        line_parts.append(raw)
                        bbox = span.get("bbox", (0.0, 0.0, 0.0, 0.0))
                        spans.append(
                            TextSpan(
                                text=raw,
                                page=page_number,
                                font=str(span.get("font", "")),
                                size=float(span.get("size", 0.0)),
                                color=int(span.get("color", -1)),
                                flags=int(span.get("flags", 0)),
                                bbox=tuple(float(v) for v in bbox),  # type: ignore[arg-type]
                            )
                        )
                    if line_parts:
                        page_lines.append("".join(line_parts))
            page_texts.append("\n".join(page_lines))

    return ParsedDocument(
        path=path,
        text=_tidy("\n\n".join(page_texts)),
        source_format="pdf",
        spans=spans,
        page_count=len(page_texts),
        page_backgrounds=backgrounds,
    )


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------


def _docx_run_colour(run) -> int:
    """Extract a run's font colour as packed 0xRRGGBB, or -1 if unset."""
    try:
        colour = run.font.color
        if colour is None or colour.rgb is None:
            return -1
        rgb = colour.rgb
        return (int(rgb[0]) << 16) | (int(rgb[1]) << 8) | int(rgb[2])
    except (AttributeError, IndexError, TypeError, ValueError):
        return -1


def _parse_docx(data: bytes, path: str) -> ParsedDocument:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ParsingError("python-docx is required to read .docx files") from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ParsingError(f"could not open DOCX {path}: {exc}") from exc

    spans: list[TextSpan] = []
    lines: list[str] = []

    def consume_paragraph(paragraph) -> None:
        parts: list[str] = []
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            parts.append(run.text)
            size = run.font.size
            spans.append(
                TextSpan(
                    text=run.text,
                    page=0,
                    font=str(run.font.name or ""),
                    size=float(size.pt) if size is not None else 0.0,
                    color=_docx_run_colour(run),
                    flags=0,
                )
            )
        if parts:
            lines.append("".join(parts))

    for paragraph in document.paragraphs:
        consume_paragraph(paragraph)

    # Resumes very often put skills and dates inside tables, so a parser that
    # only walks `document.paragraphs` silently loses them.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    consume_paragraph(paragraph)

    return ParsedDocument(
        path=path,
        text=_tidy("\n".join(lines)),
        source_format="docx",
        spans=spans,
        page_count=1,
        page_backgrounds=[0xFFFFFF],
    )


# --------------------------------------------------------------------------
# Plain text
# --------------------------------------------------------------------------


def _parse_text(data: bytes, path: str) -> ParsedDocument:
    text = data.decode("utf-8", errors="replace")
    spans = [
        TextSpan(text=line, page=0, size=11.0, color=0x000000)
        for line in text.splitlines()
        if line.strip()
    ]
    return ParsedDocument(
        path=path,
        text=_tidy(text),
        source_format="txt",
        spans=spans,
        page_count=1,
        page_backgrounds=[0xFFFFFF],
    )


# --------------------------------------------------------------------------
# Public entry points
# --------------------------------------------------------------------------

_PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".txt": _parse_text,
    ".md": _parse_text,
}


def parse_bytes(data: bytes, filename: str) -> ParsedDocument:
    """Parse an in-memory document. Used by the Streamlit uploader."""
    suffix = Path(filename).suffix.lower()
    parser = _PARSERS.get(suffix)
    if parser is None:
        raise ParsingError(
            f"unsupported file type {suffix!r} for {filename}; "
            f"supported: {', '.join(sorted(SUPPORTED_SUFFIXES))}"
        )
    return parser(data, filename)


def parse_file(path: str | Path) -> ParsedDocument:
    """Parse a document from disk."""
    path = Path(path)
    if not path.exists():
        raise ParsingError(f"file not found: {path}")
    return parse_bytes(path.read_bytes(), str(path))


def looks_scanned(document: ParsedDocument, min_chars_per_page: int = 120) -> bool:
    """Heuristic: did we get so little text that this is probably a scan?

    Image-only PDFs parse without error but yield near-empty text, which would
    otherwise surface as a candidate who simply scores zero on everything. The
    UI uses this to warn that the file needs OCR rather than silently ranking
    the person last.
    """
    if document.source_format != "pdf":
        return False
    return len(document.text) < min_chars_per_page * max(1, document.page_count)
